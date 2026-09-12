from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from src.backend.exceptions import DocumentLoadError, UnsupportedFileTypeError
from src.backend.loaders import (
    PDF_ENGINES,
    ingest_documents,
    load_docx,
    load_file,
    load_markdown,
    load_pdf,
    load_txt,
    retry_failed_documents,
)


def _create_sample_pdf(path: Path) -> None:
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text((72, 72), "Co-Work M1 first page")
    second_page = pdf.new_page()
    second_page.insert_text((72, 72), "Co-Work M1 second page")
    pdf.save(str(path))
    pdf.close()


def test_pdf_loader_supports_all_required_engines(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _create_sample_pdf(path)

    for engine in PDF_ENGINES:
        document = load_pdf(path, engine=engine)

        assert document.file_name == "sample.pdf"
        assert document.file_type == "pdf"
        assert document.document_id
        assert document.metadata["source_path"] == str(path)
        assert document.metadata["page_count"] == 2
        assert document.metadata["pages"][0]["page_number"] == 1
        assert document.metadata["pages"][1]["page_number"] == 2
        assert "Co-Work M1" in document.text


def test_pdf_loader_rejects_unknown_engine(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _create_sample_pdf(path)

    with pytest.raises(DocumentLoadError, match="Unsupported PDF engine"):
        load_pdf(path, engine="unknown")


def test_docx_loader_returns_paragraph_text_and_metadata(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    source = DocxDocument()
    source.core_properties.title = "M1 sample"
    source.add_paragraph("First paragraph")
    source.add_paragraph("Second paragraph")
    source.save(str(path))

    document = load_docx(path)

    assert document.file_name == "sample.docx"
    assert document.file_type == "docx"
    assert document.text == "First paragraph\nSecond paragraph"
    assert document.metadata["title"] == "M1 sample"
    assert document.metadata["paragraph_count"] == 2
    assert document.metadata["source_path"] == str(path)


def test_txt_and_markdown_load_as_utf8(tmp_path: Path) -> None:
    txt_path = tmp_path / "notes.txt"
    md_path = tmp_path / "notes.md"
    txt_path.write_text("中文文本", encoding="utf-8")
    md_path.write_text("# Markdown\n\n正文", encoding="utf-8")

    txt_document = load_txt(txt_path)
    md_document = load_markdown(md_path)

    assert txt_document.text == "中文文本"
    assert txt_document.file_type == "txt"
    assert txt_document.metadata["encoding"] == "utf-8"
    assert md_document.text.startswith("# Markdown")
    assert md_document.file_type == "md"


def test_text_loader_reports_invalid_utf8(tmp_path: Path) -> None:
    path = tmp_path / "invalid.txt"
    path.write_bytes(b"invalid utf-8: \xff")

    with pytest.raises(DocumentLoadError, match="UTF-8"):
        load_txt(path)


def test_load_file_dispatch_and_unsupported_type(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("dispatch", encoding="utf-8")
    assert load_file(path).text == "dispatch"

    csv_path = tmp_path / "unsupported.csv"
    csv_path.write_text("a,b\n1,2", encoding="utf-8")
    with pytest.raises(UnsupportedFileTypeError):
        load_file(csv_path)


def test_batch_ingestion_isolates_failures_and_supports_retry(tmp_path: Path) -> None:
    success_path = tmp_path / "success.txt"
    retry_path = tmp_path / "retry.txt"
    success_path.write_text("success", encoding="utf-8")
    retry_path.write_bytes(b"bad utf-8: \xff")

    result = ingest_documents([success_path, retry_path])

    assert result.successful_files == [str(success_path)]
    assert result.failed_files == [str(retry_path)]
    assert result.items[0].status == "success"
    assert result.items[1].status == "failed"
    assert result.items[1].attempts == 1
    assert "UTF-8" in result.errors[str(retry_path)]
    assert result.processed_chunk_count == 0
    assert result.index_version is None

    retry_path.write_text("fixed", encoding="utf-8")
    retried = retry_failed_documents(result)

    assert retried.failed_files == []
    assert retried.successful_files == [str(success_path), str(retry_path)]
    assert retried.items[1].status == "success"
    assert retried.items[1].attempts == 2
    assert retried.items[1].document is not None
    assert retried.items[1].document.text == "fixed"

