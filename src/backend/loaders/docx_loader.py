"""DOCX document loading using python-docx."""

from typing import Any, Dict

from ..exceptions import DocumentLoadError
from ..schemas import Document
from .base import PathLike, document_id_for, ensure_file


def load_docx(path: PathLike) -> Document:
    """Load paragraph text and core metadata from a DOCX file."""

    file_path = ensure_file(path)
    if file_path.suffix.lower() != ".docx":
        raise DocumentLoadError("Expected a .docx file: " + str(file_path))

    try:
        from docx import Document as DocxDocument

        source_document = DocxDocument(str(file_path))
    except ImportError as exc:
        raise DocumentLoadError("python-docx is required for DOCX loading") from exc
    except Exception as exc:
        raise DocumentLoadError("python-docx failed to parse " + str(file_path)) from exc

    paragraphs = [paragraph.text for paragraph in source_document.paragraphs]
    core_title = source_document.core_properties.title
    metadata: Dict[str, Any] = {
        "source_path": str(file_path),
        "title": core_title or file_path.stem,
        "paragraph_count": len(paragraphs),
    }

    return Document(
        document_id=document_id_for(file_path),
        file_name=file_path.name,
        file_type="docx",
        text="\n".join(paragraphs),
        metadata=metadata,
    )


__all__ = ["load_docx"]

